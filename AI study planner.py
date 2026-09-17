#AI powered study planner lol(⌐■_■)
#input
import tkinter as tk
from tkcalendar import DateEntry
from datetime import timedelta
from docx import Document
from docx.shared import Pt

class StudyPlannerUI:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("AI Study Planner (⌐■_■)")
        self.root.geometry("1200x720")
        self.root.config(bg="#f0f0f0")

        self.build_ui()
        self.root.mainloop()

    def build_ui(self):
        heading = tk.Label(self.root, text="🧠 AI Powered Study Planner", font=("Helvetica", 24, "bold"), bg="#f0f0f0")
        heading.pack(pady=20)

        form_frame = tk.Frame(self.root, bg="#f0f0f0")
        form_frame.pack(pady=10)

        # Subject Entry
        tk.Label(form_frame, text="Subjects (comma separated):", font=("Arial", 14), bg="#f0f0f0").grid(row=0, column=0, sticky="w")
        self.subject_entry = tk.Entry(form_frame, width=50, font=("Arial", 12))
        self.subject_entry.grid(row=0, column=1, pady=5, padx=10)

        # Topics Entry
        tk.Label(form_frame, text="Topics (one per line):", font=("Arial", 14), bg="#f0f0f0").grid(row=1, column=0, sticky="nw")
        self.topic_text = tk.Text(form_frame, height=10, width=50, font=("Arial", 12))
        self.topic_text.grid(row=1, column=1, pady=5, padx=10)

        # Date Range
        tk.Label(form_frame, text="Start Date:", font=("Arial", 14), bg="#f0f0f0").grid(row=2, column=0, sticky="w")
        self.start_date = DateEntry(form_frame, width=20, font=("Arial", 12))
        self.start_date.grid(row=2, column=1, sticky="w", pady=5)

        tk.Label(form_frame, text="End Date:", font=("Arial", 14), bg="#f0f0f0").grid(row=3, column=0, sticky="w")
        self.end_date = DateEntry(form_frame, width=20, font=("Arial", 12))
        self.end_date.grid(row=3, column=1, sticky="w", pady=5)

        # Summer Break Dates
        tk.Label(form_frame, text="Summer Break Start:", font=("Arial", 14), bg="#f0f0f0").grid(row=4, column=0, sticky="w")
        self.summer_start = DateEntry(form_frame, width=20, font=("Arial", 12))
        self.summer_start.grid(row=4, column=1, sticky="w", pady=5)

        tk.Label(form_frame, text="Summer Break End:", font=("Arial", 14), bg="#f0f0f0").grid(row=5, column=0, sticky="w")
        self.summer_end = DateEntry(form_frame, width=20, font=("Arial", 12))
        self.summer_end.grid(row=5, column=1, sticky="w", pady=5)

        # Submit Button
        self.submit_btn = tk.Button(self.root, text="Generate Plan", font=("Arial", 16, "bold"), bg="#4CAF50", fg="white", command=self.get_inputs)
        self.submit_btn.pack(pady=20)

        # Output label
        self.output_label = tk.Label(self.root, text="", font=("Arial", 12), fg="green", bg="#f0f0f0")
        self.output_label.pack()

        def get_inputs(self):
            subjects = self.subject_entry.get().strip()
            topics = self.topic_text.get("1.0", tk.END).strip().splitlines()
            start = self.start_date.get_date()
            end = self.end_date.get_date()
            summer_start = self.summer_start.get_date()
            summer_end = self.summer_end.get_date()

            if not topics or start >= end:
                self.output_label.config(text="❌ Invalid input. Check topic list or date range.", fg="red")
            return

    # 1. Generate the study plan
            plan = self.generate_study_plan(topics, start, end, summer_start, summer_end)

    # 2. Export it
            self.export_to_word(plan)
            self.output_label.config(text="✅ Word document created: Study_Plan.docx", fg="green")


        ## TODO: Plug this into your logic & word doc generation
    def generate_study_plan(self, topics, start, end, summer_start, summer_end):
        day = start
        plan = []
        idx = 0

        while day <= end and idx < len(topics):
        # Determine daily study hours
            if summer_start <= day <= summer_end or day.weekday() == 6:
                hours = 12
            else:
                hours = 6

        todays_topics = topics[idx:idx + hours]
        plan.append((day.strftime("%d-%b-%Y"), day.strftime("%A"), todays_topics))
        idx += hours
        day += timedelta(days=1)

        return plan

    def export_to_word(self, plan, filename="Study_Plan.docx"):
        doc = Document()
        doc.add_heading("🧠 Study Plan", level=1)

        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Date'
        hdr_cells[1].text = 'Day'
        hdr_cells[2].text = 'Topics'

        for date, day, topics in plan:
            row_cells = table.add_row().cells
            row_cells[0].text = date
            row_cells[1].text = day
            row_cells[2].text = ', '.join(topics)

    # Formatting font size
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)

        doc.save(filename)
        

# Run the UI
StudyPlannerUI()
